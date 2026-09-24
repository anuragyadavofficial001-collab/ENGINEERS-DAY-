# ============================================================
# ENGINEERS DAY 2026 - ADMIN CONTROL CENTER
# PostgreSQL / Supabase Compatible
# ============================================================

import csv
import io
import os
import re
from functools import wraps
from io import BytesIO

from flask import (
    Blueprint,
    current_app,
    flash,
    redirect,
    render_template,
    request,
    session,
    send_file,
    url_for
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from werkzeug.security import check_password_hash

from database.db import get_db_connection


# ============================================================
# BLUEPRINT
# ============================================================

admin_bp = Blueprint(
    "admin",
    __name__,
    url_prefix="/admin"
)


# ============================================================
# CONSTANTS
# ============================================================

ALLOWED_EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".pdf"
}

MAX_FILE_SIZE = 10 * 1024 * 1024

ALLOWED_REGISTRATION_MODES = {
    "INDIVIDUAL",
    "TEAM"
}

ALLOWED_GAME_STATUSES = {
    "UPCOMING",
    "REGISTRATION_OPEN",
    "REGISTRATION_CLOSED",
    "LIVE",
    "COMPLETED",
    "CANCELLED"
}

ALLOWED_REGISTRATION_STATUSES = {
    "REGISTERED",
    "CANCELLED",
    "APPROVED",
    "REJECTED"
}

ALLOWED_NOTIFICATION_TYPES = {
    "GENERAL",
    "EVENT",
    "RESULT",
    "IMPORTANT"
}

ALLOWED_PORTAL_STATUSES = {
    "ACTIVE",
    "MAINTENANCE"
}


# ============================================================
# SMALL HELPERS
# ============================================================

def close_db(connection=None, cursor=None):
    """
    Safely close cursor and database connection.
    """

    try:
        if cursor:
            cursor.close()
    except Exception:
        pass

    try:
        if connection:
            connection.close()
    except Exception:
        pass


def rollback_db(connection=None):
    """
    Safely rollback current transaction.
    """

    try:
        if connection:
            connection.rollback()
    except Exception:
        pass


def clean_text(value):
    """
    Normalize text values.
    """

    if value is None:
        return None

    value = str(value).strip()

    return value if value else None


def clean_int(value, default=None):
    """
    Convert a value to integer safely.
    """

    if value is None:
        return default

    value = str(value).strip()

    if not value:
        return default

    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def form_checkbox(name):
    """
    Convert an HTML checkbox into a Python boolean.

    PostgreSQL boolean columns should receive True/False,
    not MySQL-style 1/0.
    """

    return bool(request.form.get(name))


# ============================================================
# ADMIN AUTH
# ============================================================

def admin_required(view_function):

    @wraps(view_function)
    def wrapped_view(*args, **kwargs):

        if not session.get("admin_authenticated"):
            return redirect(
                url_for("admin.login")
            )

        return view_function(*args, **kwargs)

    return wrapped_view


# ============================================================
# ADMIN LOGIN
# ============================================================

@admin_bp.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    if session.get("admin_authenticated"):
        return redirect(
            url_for("admin.dashboard")
        )

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        if not username or not password:

            flash(
                "Username and password are required.",
                "error"
            )

            return render_template(
                "admin/login.html"
            )

        connection = None
        cursor = None

        try:

            connection = get_db_connection()
            cursor = connection.cursor()

            cursor.execute(
                """
                SELECT
                    id,
                    username,
                    password_hash,
                    is_active
                FROM admins
                WHERE username = %s
                LIMIT 1
                """,
                (username,)
            )

            admin = cursor.fetchone()

            if not admin:

                flash(
                    "Invalid admin credentials.",
                    "error"
                )

                return render_template(
                    "admin/login.html"
                )

            if not bool(admin["is_active"]):

                flash(
                    "This admin account is disabled.",
                    "error"
                )

                return render_template(
                    "admin/login.html"
                )

            try:
                password_valid = check_password_hash(
                    admin["password_hash"],
                    password
                )
            except Exception:

                current_app.logger.exception(
                    "Invalid admin password hash"
                )

                password_valid = False

            if not password_valid:

                flash(
                    "Invalid admin credentials.",
                    "error"
                )

                return render_template(
                    "admin/login.html"
                )

            # ------------------------------------------------
            # New clean admin session
            # ------------------------------------------------

            session.clear()

            session["admin_authenticated"] = True
            session["admin_id"] = admin["id"]
            session["admin_username"] = admin["username"]
            session.permanent = True

            return redirect(
                url_for("admin.dashboard")
            )

        except Exception:

            rollback_db(connection)

            current_app.logger.exception(
                "Admin login error"
            )

            flash(
                "Unable to process login right now.",
                "error"
            )

            return render_template(
                "admin/login.html"
            )

        finally:

            close_db(
                connection,
                cursor
            )

    return render_template(
        "admin/login.html"
    )


# ============================================================
# ADMIN LOGOUT
# ============================================================

@admin_bp.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("admin.login")
    )


# ============================================================
# ADMIN DASHBOARD
# ============================================================

@admin_bp.route("/dashboard")
@admin_required
def dashboard():

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        # ----------------------------------------------------
        # TOTAL ACTIVE STUDENT REGISTRY
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM student_registry
            WHERE is_active = TRUE
            """
        )

        total_students = cursor.fetchone()["total"]

        # ----------------------------------------------------
        # ACTIVE STUDENT ACCOUNTS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM students
            WHERE is_active = TRUE
            """
        )

        active_accounts = cursor.fetchone()["total"]

        # ----------------------------------------------------
        # TOTAL EVENTS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM games
            """
        )

        total_events = cursor.fetchone()["total"]

        # ----------------------------------------------------
        # TOTAL REGISTRATIONS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM registrations
            """
        )

        total_registrations = cursor.fetchone()["total"]

        return render_template(
            "admin/dashboard.html",
            total_students=total_students,
            active_accounts=active_accounts,
            total_events=total_events,
            total_registrations=total_registrations
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Admin dashboard error"
        )

        flash(
            "Unable to load admin dashboard.",
            "error"
        )

        return redirect(
            url_for("admin.login")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# STUDENT REGISTRY
#
# Search:
# ?q=student_id
# ?q=name
# ?q=email
# ?q=phone
#
# Filters:
# ?status=all
# ?status=active
# ?status=disabled
# ?status=claimed
# ?status=unclaimed
# ============================================================

@admin_bp.route("/students")
@admin_required
def students():

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        search = request.args.get(
            "search",
            request.args.get("q", "")
        ).strip()

        status = request.args.get(
            "status",
            "all"
        ).strip().lower()

        # ----------------------------------------------------
        # BASE QUERY
        # ----------------------------------------------------

        query = """
            SELECT
                sr.id AS registry_id,
                sr.student_id,
                sr.is_active AS registry_active,
                sr.claimed,
                sr.claimed_student_id,
                sr.imported_at,
                sr.claimed_at,

                s.id AS account_id,
                s.name,
                s.email,
                s.phone,
                s.branch,
                s.section,
                s.year,
                s.is_active AS account_active

            FROM student_registry sr

            LEFT JOIN students s
                ON (
                    s.id = sr.claimed_student_id
                    OR s.student_id = sr.student_id
                )
        """

        conditions = []
        params = []

        # ----------------------------------------------------
        # SEARCH
        # PostgreSQL uses ILIKE for case-insensitive search.
        # ----------------------------------------------------

        if search:

            conditions.append(
                """
                (
                    sr.student_id ILIKE %s
                    OR s.student_id ILIKE %s
                    OR s.name ILIKE %s
                    OR s.email ILIKE %s
                    OR s.phone ILIKE %s
                    OR s.branch ILIKE %s
                    OR s.section ILIKE %s
                )
                """
            )

            search_value = f"%{search}%"

            params.extend([
                search_value,
                search_value,
                search_value,
                search_value,
                search_value,
                search_value,
                search_value
            ])

        # ----------------------------------------------------
        # STATUS
        # ----------------------------------------------------

        if status == "active":

            conditions.append(
                "sr.is_active = TRUE"
            )

        elif status == "disabled":

            conditions.append(
                "sr.is_active = FALSE"
            )

        elif status == "claimed":

            conditions.append(
                "sr.claimed = TRUE"
            )

        elif status == "unclaimed":

            conditions.append(
                "sr.claimed = FALSE"
            )

        if conditions:

            query += (
                " WHERE "
                + " AND ".join(conditions)
            )

        query += """
            ORDER BY sr.id DESC
            LIMIT 1000
        """

        cursor.execute(
            query,
            tuple(params)
        )

        registry = cursor.fetchall()

        # ----------------------------------------------------
        # STATISTICS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM student_registry
            """
        )

        total = cursor.fetchone()["total"]

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM student_registry
            WHERE is_active = TRUE
            """
        )

        active = cursor.fetchone()["total"]

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM student_registry
            WHERE claimed = TRUE
            """
        )

        claimed = cursor.fetchone()["total"]

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM student_registry
            WHERE claimed = FALSE
            """
        )

        unclaimed = cursor.fetchone()["total"]

        return render_template(
            "admin/students.html",
            registry=registry,
            total=total,
            active=active,
            claimed=claimed,
            unclaimed=unclaimed,
            search=search,
            status=status
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Student registry error"
        )

        flash(
            "Unable to load student registry.",
            "error"
        )

        return redirect(
            url_for("admin.dashboard")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# STUDENT DETAIL
# ============================================================

@admin_bp.route(
    "/students/<int:registry_id>/view"
)
@admin_required
def student_detail(registry_id):

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT
                sr.id AS registry_id,
                sr.student_id,
                sr.is_active AS registry_active,
                sr.claimed,
                sr.claimed_student_id,
                sr.imported_at,
                sr.claimed_at,

                s.id AS account_id,
                s.student_id AS account_student_id,
                s.name,
                s.email,
                s.phone,
                s.branch,
                s.section,
                s.year,
                s.is_active AS account_active

            FROM student_registry sr

            LEFT JOIN students s
                ON (
                    s.id = sr.claimed_student_id
                    OR s.student_id = sr.student_id
                )

            WHERE sr.id = %s

            LIMIT 1
            """,
            (registry_id,)
        )

        student = cursor.fetchone()

        if not student:

            flash(
                "Student record not found.",
                "error"
            )

            return redirect(
                url_for("admin.students")
            )

        return render_template(
            "admin/student_detail.html",
            student=student
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Student detail error"
        )

        flash(
            "Unable to load student details.",
            "error"
        )

        return redirect(
            url_for("admin.students")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# NORMALIZE STUDENT ID
# ============================================================

def normalize_student_id(value):

    if value is None:
        return None

    value = str(value).strip()

    if not value:
        return None

    value = value.upper()

    if not re.fullmatch(
        r"[A-Z0-9][A-Z0-9._/-]{0,49}",
        value
    ):
        return None

    return value


# ============================================================
# NORMALIZE PHONE
# ============================================================

def normalize_phone(value):

    value = clean_text(value)

    if not value:
        return None

    value = re.sub(
        r"[\s()-]",
        "",
        value
    )

    if value.startswith("+91"):
        value = value[3:]

    if value.startswith("91") and len(value) == 12:
        value = value[2:]

    if re.fullmatch(
        r"[6-9]\d{9}",
        value
    ):
        return value

    return None


# ============================================================
# NORMALIZE HEADER
# ============================================================

def normalize_header(value):

    if value is None:
        return ""

    value = str(value).strip().lower()

    return re.sub(
        r"[^a-z0-9]",
        "",
        value
    )


# ============================================================
# FIND COLUMN
# ============================================================

def find_column(
    headers,
    possible_names
):

    normalized = {}

    for index, header in enumerate(headers):

        normalized[
            normalize_header(header)
        ] = index

    for name in possible_names:

        key = normalize_header(name)

        if key in normalized:
            return normalized[key]

    return None


# ============================================================
# COLUMN ALIASES
# ============================================================

ID_COLUMNS = [
    "Student ID",
    "StudentID",
    "Roll Number",
    "Roll No",
    "Roll",
    "Enrollment Number",
    "Enrollment No",
    "Enrollment"
]

NAME_COLUMNS = [
    "Name",
    "Student Name",
    "Full Name"
]

EMAIL_COLUMNS = [
    "Email",
    "Email ID",
    "Email Address"
]

PHONE_COLUMNS = [
    "Phone",
    "Mobile",
    "Mobile Number",
    "Phone Number",
    "Contact"
]

BRANCH_COLUMNS = [
    "Branch",
    "Course",
    "Program",
    "Department"
]

SECTION_COLUMNS = [
    "Section",
    "Sec"
]

YEAR_COLUMNS = [
    "Year",
    "Academic Year",
    "Year of Study"
]


# ============================================================
# BUILD STUDENT RECORD
# ============================================================

def build_record(
    row,
    headers
):

    id_column = find_column(
        headers,
        ID_COLUMNS
    )

    if id_column is None:
        return None

    student_id = normalize_student_id(
        row[id_column]
        if id_column < len(row)
        else None
    )

    if not student_id:
        return None

    name_column = find_column(
        headers,
        NAME_COLUMNS
    )

    email_column = find_column(
        headers,
        EMAIL_COLUMNS
    )

    phone_column = find_column(
        headers,
        PHONE_COLUMNS
    )

    branch_column = find_column(
        headers,
        BRANCH_COLUMNS
    )

    section_column = find_column(
        headers,
        SECTION_COLUMNS
    )

    year_column = find_column(
        headers,
        YEAR_COLUMNS
    )

    def get_value(column):

        if column is None:
            return None

        if column >= len(row):
            return None

        return clean_text(
            row[column]
        )

    phone = normalize_phone(
        get_value(phone_column)
    )

    return {
        "student_id": student_id,
        "name": get_value(name_column),
        "email": get_value(email_column),
        "phone": phone,
        "branch": get_value(branch_column),
        "section": get_value(section_column),
        "year": get_value(year_column)
    }


# ============================================================
# EXTRACT CSV
# ============================================================

def extract_from_csv(file_bytes):

    text = file_bytes.decode(
        "utf-8-sig",
        errors="ignore"
    )

    reader = csv.reader(
        io.StringIO(text)
    )

    rows = list(reader)

    if not rows:
        return []

    headers = rows[0]

    if find_column(
        headers,
        ID_COLUMNS
    ) is None:

        raise ValueError(
            "Student ID / Roll Number column not found."
        )

    results = []

    for row in rows[1:]:

        record = build_record(
            row,
            headers
        )

        if record:
            results.append(record)

    return results


# ============================================================
# EXTRACT XLSX
# ============================================================

def extract_from_excel(file_bytes):

    workbook = load_workbook(
        io.BytesIO(file_bytes),
        read_only=True,
        data_only=True
    )

    results = []

    try:

        for sheet in workbook.worksheets:

            rows = sheet.iter_rows(
                values_only=True
            )

            try:
                headers = list(next(rows))
            except StopIteration:
                continue

            if find_column(
                headers,
                ID_COLUMNS
            ) is None:

                continue

            for row in rows:

                row = list(row)

                record = build_record(
                    row,
                    headers
                )

                if record:
                    results.append(record)

    finally:

        workbook.close()

    if not results:

        raise ValueError(
            "No valid student records were found."
        )

    return results


# ============================================================
# EXTRACT PDF
#
# PDF is ID-only.
# ============================================================

def extract_from_pdf(file_bytes):

    reader = PdfReader(
        io.BytesIO(file_bytes)
    )

    results = []

    for page in reader.pages:

        text = page.extract_text() or ""

        candidates = re.findall(
            r"\b[A-Z0-9][A-Z0-9._/-]{2,49}\b",
            text.upper()
        )

        for candidate in candidates:

            student_id = normalize_student_id(
                candidate
            )

            if student_id:

                results.append({
                    "student_id": student_id,
                    "name": None,
                    "email": None,
                    "phone": None,
                    "branch": None,
                    "section": None,
                    "year": None
                })

    return results


# ============================================================
# MAIN EXTRACTION
# ============================================================

def extract_student_records(
    filename,
    file_bytes
):

    extension = os.path.splitext(
        filename
    )[1].lower()

    if extension not in ALLOWED_EXTENSIONS:

        raise ValueError(
            "Unsupported file format. "
            "Only CSV, XLSX and PDF are allowed."
        )

    if len(file_bytes) > MAX_FILE_SIZE:

        raise ValueError(
            "File size cannot exceed 10 MB."
        )

    if extension == ".csv":

        records = extract_from_csv(
            file_bytes
        )

    elif extension == ".xlsx":

        records = extract_from_excel(
            file_bytes
        )

    elif extension == ".pdf":

        records = extract_from_pdf(
            file_bytes
        )

    else:

        records = []

    # --------------------------------------------------------
    # REMOVE DUPLICATE STUDENT IDs
    # --------------------------------------------------------

    unique = {}

    for record in records:

        student_id = record["student_id"]

        if student_id not in unique:

            unique[student_id] = record

        else:

            old = unique[student_id]

            for field in [
                "name",
                "email",
                "phone",
                "branch",
                "section",
                "year"
            ]:

                if (
                    not old.get(field)
                    and record.get(field)
                ):

                    old[field] = record[field]

    return list(
        unique.values()
    )


# ============================================================
# SAVE STUDENT RECORD
#
# Registry = authorization
# Students  = profile/account
#
# PostgreSQL FIX:
# INSERT ... RETURNING id
# ============================================================

def save_student_record(
    cursor,
    record
):

    student_id = record["student_id"]

    # --------------------------------------------------------
    # CHECK REGISTRY
    # --------------------------------------------------------

    cursor.execute(
        """
        SELECT
            id,
            claimed,
            is_active
        FROM student_registry
        WHERE student_id = %s
        LIMIT 1
        """,
        (student_id,)
    )

    registry = cursor.fetchone()

    if registry:

        registry_id = registry["id"]

    else:

        cursor.execute(
            """
            INSERT INTO student_registry
            (
                student_id,
                is_active,
                claimed
            )
            VALUES
            (
                %s,
                TRUE,
                FALSE
            )
            RETURNING id
            """,
            (student_id,)
        )

        inserted_registry = cursor.fetchone()

        if not inserted_registry:

            raise RuntimeError(
                "Unable to create student registry record."
            )

        registry_id = inserted_registry["id"]

    # --------------------------------------------------------
    # CHECK STUDENT PROFILE
    # --------------------------------------------------------

    cursor.execute(
        """
        SELECT
            id
        FROM students
        WHERE student_id = %s
        LIMIT 1
        """,
        (student_id,)
    )

    existing_student = cursor.fetchone()

    fields = {
        "name": record.get("name"),
        "email": record.get("email"),
        "phone": record.get("phone"),
        "branch": record.get("branch"),
        "section": record.get("section"),
        "year": record.get("year")
    }

    # --------------------------------------------------------
    # UPDATE EXISTING PROFILE
    # --------------------------------------------------------

    if existing_student:

        updates = []
        values = []

        for field, value in fields.items():

            if value is not None:

                updates.append(
                    f"{field} = %s"
                )

                values.append(value)

        if updates:

            values.append(
                existing_student["id"]
            )

            cursor.execute(
                f"""
                UPDATE students
                SET {", ".join(updates)}
                WHERE id = %s
                """,
                tuple(values)
            )

        return "updated"

    # --------------------------------------------------------
    # CREATE NEW PROFILE
    # --------------------------------------------------------

    cursor.execute(
        """
        INSERT INTO students
        (
            name,
            email,
            phone,
            student_id,
            branch,
            section,
            year,
            is_active
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            TRUE
        )
        """,
        (
            fields["name"] or "Pending Student",
            fields["email"],
            fields["phone"],
            student_id,
            fields["branch"],
            fields["section"],
            fields["year"]
        )
    )

    return "added"


# ============================================================
# UPLOAD STUDENTS
# ============================================================

@admin_bp.route(
    "/students/upload",
    methods=["GET", "POST"]
)
@admin_required
def upload_students():

    if request.method == "GET":

        return render_template(
            "admin/upload_students.html"
        )

    uploaded_file = request.files.get(
        "student_file"
    )

    if not uploaded_file:

        flash(
            "Please select a file.",
            "error"
        )

        return redirect(
            url_for("admin.upload_students")
        )

    filename = (
        uploaded_file.filename or ""
    ).strip()

    if not filename:

        flash(
            "Invalid file.",
            "error"
        )

        return redirect(
            url_for("admin.upload_students")
        )

    file_bytes = uploaded_file.read()

    if not file_bytes:

        flash(
            "The uploaded file is empty.",
            "error"
        )

        return redirect(
            url_for("admin.upload_students")
        )

    connection = None
    cursor = None

    try:

        records = extract_student_records(
            filename,
            file_bytes
        )

        if not records:

            flash(
                "No valid student records were found.",
                "error"
            )

            return redirect(
                url_for("admin.upload_students")
            )

        connection = get_db_connection()
        cursor = connection.cursor()

        added = 0
        updated = 0

        for record in records:

            result = save_student_record(
                cursor,
                record
            )

            if result == "added":
                added += 1

            elif result == "updated":
                updated += 1

        connection.commit()

        flash(
            f"Import successful — "
            f"{added} new students, "
            f"{updated} existing profiles updated.",
            "success"
        )

        return redirect(
            url_for("admin.students")
        )

    except ValueError as error:

        rollback_db(connection)

        flash(
            str(error),
            "error"
        )

        return redirect(
            url_for("admin.upload_students")
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Student import error"
        )

        flash(
            "Student import failed. "
            "Please check the file format and data.",
            "error"
        )

        return redirect(
            url_for("admin.upload_students")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# MANUAL STUDENT ID IMPORT
# ============================================================

@admin_bp.route(
    "/students/import",
    methods=["POST"]
)
@admin_required
def import_students():

    raw_ids = request.form.get(
        "student_ids",
        ""
    )

    if not raw_ids:

        flash(
            "No student data available.",
            "error"
        )

        return redirect(
            url_for("admin.upload_students")
        )

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        added = 0
        duplicate = 0

        for value in raw_ids.splitlines():

            student_id = normalize_student_id(
                value
            )

            if not student_id:
                continue

            cursor.execute(
                """
                SELECT id
                FROM student_registry
                WHERE student_id = %s
                LIMIT 1
                """,
                (student_id,)
            )

            existing = cursor.fetchone()

            if existing:

                duplicate += 1
                continue

            cursor.execute(
                """
                INSERT INTO student_registry
                (
                    student_id,
                    is_active,
                    claimed
                )
                VALUES
                (
                    %s,
                    TRUE,
                    FALSE
                )
                """,
                (student_id,)
            )

            added += 1

        connection.commit()

        flash(
            f"Student IDs imported — "
            f"{added} added, "
            f"{duplicate} duplicates skipped.",
            "success"
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Manual student import error"
        )

        flash(
            "Student import failed.",
            "error"
        )

    finally:

        close_db(
            connection,
            cursor
        )

    return redirect(
        url_for("admin.students")
    )


# ============================================================
# TOGGLE STUDENT
#
# Registry and student account remain synchronized.
# ============================================================

@admin_bp.route(
    "/students/<int:registry_id>/toggle",
    methods=["POST"]
)
@admin_required
def toggle_student(registry_id):

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT
                sr.id,
                sr.student_id,
                sr.is_active,
                sr.claimed_student_id,
                s.id AS account_id
            FROM student_registry sr

            LEFT JOIN students s
                ON (
                    s.id = sr.claimed_student_id
                    OR s.student_id = sr.student_id
                )

            WHERE sr.id = %s

            LIMIT 1
            """,
            (registry_id,)
        )

        student = cursor.fetchone()

        if not student:

            flash(
                "Student record not found.",
                "error"
            )

            return redirect(
                url_for("admin.students")
            )

        # ----------------------------------------------------
        # PostgreSQL boolean handling
        # ----------------------------------------------------

        new_status = not bool(
            student["is_active"]
        )

        # ----------------------------------------------------
        # UPDATE REGISTRY
        # ----------------------------------------------------

        cursor.execute(
            """
            UPDATE student_registry
            SET is_active = %s
            WHERE id = %s
            """,
            (
                new_status,
                registry_id
            )
        )

        # ----------------------------------------------------
        # UPDATE ACTUAL STUDENT ACCOUNT
        # ----------------------------------------------------

        if student["account_id"]:

            cursor.execute(
                """
                UPDATE students
                SET is_active = %s
                WHERE id = %s
                """,
                (
                    new_status,
                    student["account_id"]
                )
            )

        connection.commit()

        if new_status:

            flash(
                f"{student['student_id']} has been enabled.",
                "success"
            )

        else:

            flash(
                f"{student['student_id']} has been disabled.",
                "success"
            )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Student toggle error"
        )

        flash(
            "Unable to update student status.",
            "error"
        )

    finally:

        close_db(
            connection,
            cursor
        )

    return redirect(
        url_for("admin.students")
    )


# ============================================================
# DELETE STUDENT
#
# Only UNCLAIMED registry records can be directly deleted.
# ============================================================

@admin_bp.route(
    "/students/<int:registry_id>/delete",
    methods=["POST"]
)
@admin_required
def delete_student(registry_id):

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT
                id,
                student_id,
                claimed,
                claimed_student_id
            FROM student_registry
            WHERE id = %s
            LIMIT 1
            """,
            (registry_id,)
        )

        student = cursor.fetchone()

        if not student:

            flash(
                "Student record not found.",
                "error"
            )

            return redirect(
                url_for("admin.students")
            )

        if bool(student["claimed"]):

            flash(
                f"{student['student_id']} is already claimed. "
                f"Disable it instead of deleting it.",
                "error"
            )

            return redirect(
                url_for("admin.students")
            )

        cursor.execute(
            """
            DELETE FROM student_registry
            WHERE id = %s
              AND claimed = FALSE
            """,
            (registry_id,)
        )

        deleted_rows = cursor.rowcount

        connection.commit()

        if deleted_rows == 0:

            flash(
                "Student record could not be deleted.",
                "error"
            )

        else:

            flash(
                f"Student ID {student['student_id']} "
                f"was deleted successfully.",
                "success"
            )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Student deletion error"
        )

        flash(
            "Unable to delete student record.",
            "error"
        )

    finally:

        close_db(
            connection,
            cursor
        )

    return redirect(
        url_for("admin.students")
    )


# ============================================================
# EVENT MANAGEMENT
# ============================================================

@admin_bp.route("/events")
@admin_required
def events():

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT
                id,
                game_name,
                description,
                registration_mode,
                team_min_size,
                team_max_size,
                event_date,
                start_time,
                end_time,

                CONCAT_WS(
                    ', ',
                    NULLIF(block, ''),
                    NULLIF(floor, ''),
                    NULLIF(room, '')
                ) AS venue,

                block,
                floor,
                room,

                prize_pool,
                rules,
                requirements,
                registration_open,
                status,
                winner_certificate,
                runner_up_certificate

            FROM games

            ORDER BY id ASC
            """
        )

        games = cursor.fetchall()

        return render_template(
            "admin/events.html",
            games=games
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Admin event management error"
        )

        flash(
            "Unable to load event management.",
            "error"
        )

        return redirect(
            url_for("admin.dashboard")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# EDIT EVENT
# ============================================================

@admin_bp.route(
    "/events/<int:game_id>/edit",
    methods=["GET", "POST"]
)
@admin_required
def edit_event(game_id):

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        # ----------------------------------------------------
        # LOAD EVENT
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT *
            FROM games
            WHERE id = %s
            LIMIT 1
            """,
            (game_id,)
        )

        game = cursor.fetchone()

        if not game:

            flash(
                "Event not found.",
                "error"
            )

            return redirect(
                url_for("admin.events")
            )

        # ----------------------------------------------------
        # POST
        # ----------------------------------------------------

        if request.method == "POST":

            game_name = request.form.get(
                "game_name",
                ""
            ).strip()

            description = request.form.get(
                "description",
                ""
            ).strip() or None

            registration_mode = request.form.get(
                "registration_mode",
                "INDIVIDUAL"
            ).strip().upper()

            team_min_size = clean_int(
                request.form.get(
                    "team_min_size"
                ),
                0
            )

            team_max_size = clean_int(
                request.form.get(
                    "team_max_size"
                ),
                0
            )

            event_date = request.form.get(
                "event_date",
                ""
            ).strip() or None

            start_time = request.form.get(
                "start_time",
                ""
            ).strip() or None

            end_time = request.form.get(
                "end_time",
                ""
            ).strip() or None

            block = request.form.get(
                "block",
                ""
            ).strip() or None

            floor = request.form.get(
                "floor",
                ""
            ).strip() or None

            room = request.form.get(
                "room",
                ""
            ).strip() or None

            prize_pool = request.form.get(
                "prize_pool",
                ""
            ).strip() or None

            rules = request.form.get(
                "rules",
                ""
            ).strip() or None

            requirements = request.form.get(
                "requirements",
                ""
            ).strip() or None

            registration_open = form_checkbox(
                "registration_open"
            )

            status = request.form.get(
                "status",
                "UPCOMING"
            ).strip().upper()

            winner_certificate = form_checkbox(
                "winner_certificate"
            )

            runner_up_certificate = form_checkbox(
                "runner_up_certificate"
            )

            # ------------------------------------------------
            # VALIDATION
            # ------------------------------------------------

            if not game_name:

                flash(
                    "Event name is required.",
                    "error"
                )

                return redirect(
                    url_for(
                        "admin.edit_event",
                        game_id=game_id
                    )
                )

            if registration_mode not in ALLOWED_REGISTRATION_MODES:

                flash(
                    "Invalid registration mode.",
                    "error"
                )

                return redirect(
                    url_for(
                        "admin.edit_event",
                        game_id=game_id
                    )
                )

            if status not in ALLOWED_GAME_STATUSES:

                flash(
                    "Invalid event status.",
                    "error"
                )

                return redirect(
                    url_for(
                        "admin.edit_event",
                        game_id=game_id
                    )
                )

            if team_min_size is None:
                team_min_size = 0

            if team_max_size is None:
                team_max_size = 0

            if registration_mode == "INDIVIDUAL":

                team_min_size = 0
                team_max_size = 0

            else:

                if team_min_size < 1:

                    flash(
                        "Team minimum size must be at least 1.",
                        "error"
                    )

                    return redirect(
                        url_for(
                            "admin.edit_event",
                            game_id=game_id
                        )
                    )

                if team_max_size < team_min_size:

                    flash(
                        "Maximum team size cannot be smaller than minimum.",
                        "error"
                    )

                    return redirect(
                        url_for(
                            "admin.edit_event",
                            game_id=game_id
                        )
                    )

            # ------------------------------------------------
            # UPDATE EVENT
            # ------------------------------------------------

            cursor.execute(
                """
                UPDATE games
                SET
                    game_name = %s,
                    description = %s,
                    registration_mode = %s,
                    team_min_size = %s,
                    team_max_size = %s,
                    event_date = %s,
                    start_time = %s,
                    end_time = %s,
                    block = %s,
                    floor = %s,
                    room = %s,
                    prize_pool = %s,
                    rules = %s,
                    requirements = %s,
                    registration_open = %s,
                    status = %s,
                    winner_certificate = %s,
                    runner_up_certificate = %s,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = %s
                """,
                (
                    game_name,
                    description,
                    registration_mode,
                    team_min_size,
                    team_max_size,
                    event_date,
                    start_time,
                    end_time,
                    block,
                    floor,
                    room,
                    prize_pool,
                    rules,
                    requirements,
                    registration_open,
                    status,
                    winner_certificate,
                    runner_up_certificate,
                    game_id
                )
            )

            connection.commit()

            flash(
                "Event updated successfully.",
                "success"
            )

            return redirect(
                url_for("admin.events")
            )

        # ----------------------------------------------------
        # GET
        # ----------------------------------------------------

        return render_template(
            "admin/event_edit.html",
            game=game
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Event update error"
        )

        flash(
            "Unable to update event.",
            "error"
        )

        return redirect(
            url_for("admin.events")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# REGISTRATION MANAGEMENT
# ============================================================

@admin_bp.route("/registrations")
@admin_required
def registrations():

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        search = request.args.get(
            "q",
            ""
        ).strip()

        status = request.args.get(
            "status",
            "all"
        ).strip().upper()

        game_id = request.args.get(
            "game_id",
            ""
        ).strip()

        query = """
            SELECT
                r.id,
                r.student_id,
                r.game_id,
                r.status,
                r.registered_at,
                r.updated_at,

                s.name AS student_name,
                s.name AS name,
                s.student_id AS student_roll,
                s.student_id AS student_id_display,
                s.email AS student_email,
                s.email AS email,
                s.phone AS student_phone,
                s.phone AS phone,
                s.branch,
                s.section,
                s.year,

                g.game_name,
                g.registration_mode,
                g.event_date,
                g.event_date AS schedule_date,
                g.start_time,

                t.id AS team_id,
                t.team_name,
                t.status AS team_status,
                tm.member_role AS team_role,

                CONCAT_WS(
                    ', ',
                    NULLIF(g.block, ''),
                    NULLIF(g.floor, ''),
                    NULLIF(g.room, '')
                ) AS venue

            FROM registrations r

            INNER JOIN students s
                ON s.id = r.student_id

            INNER JOIN games g
                ON g.id = r.game_id

            LEFT JOIN team_members tm
                ON tm.student_id = r.student_id

            LEFT JOIN teams t
                ON t.id = tm.team_id
               AND t.game_id = r.game_id
        """

        conditions = []
        params = []

        # ----------------------------------------------------
        # SEARCH
        # ----------------------------------------------------

        if search:

            value = f"%{search}%"

            conditions.append(
                """
                (
                    s.name ILIKE %s
                    OR s.student_id ILIKE %s
                    OR s.email ILIKE %s
                    OR s.phone ILIKE %s
                    OR g.game_name ILIKE %s
                )
                """
            )

            params.extend([
                value,
                value,
                value,
                value,
                value
            ])

        # ----------------------------------------------------
        # STATUS
        # ----------------------------------------------------

        if status in ALLOWED_REGISTRATION_STATUSES:

            conditions.append(
                "r.status = %s"
            )

            params.append(
                status
            )

        # ----------------------------------------------------
        # EVENT
        # ----------------------------------------------------

        if game_id.isdigit():

            conditions.append(
                "r.game_id = %s"
            )

            params.append(
                int(game_id)
            )

        if conditions:

            query += (
                " WHERE "
                + " AND ".join(conditions)
            )

        query += """
            ORDER BY
                r.registered_at DESC,
                r.id DESC
            LIMIT 2000
        """

        cursor.execute(
            query,
            tuple(params)
        )

        registration_rows = cursor.fetchall()

        # ----------------------------------------------------
        # EVENTS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT
                id,
                game_name
            FROM games
            ORDER BY game_name ASC
            """
        )

        games = cursor.fetchall()

        # ----------------------------------------------------
        # STATISTICS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM registrations
            """
        )

        total = cursor.fetchone()["total"]

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM registrations
            WHERE status = 'REGISTERED'
            """
        )

        registered = cursor.fetchone()["total"]

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM registrations
            WHERE status = 'APPROVED'
            """
        )

        approved = cursor.fetchone()["total"]

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM registrations
            WHERE status = 'REJECTED'
            """
        )

        rejected = cursor.fetchone()["total"]

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM registrations
            WHERE status = 'CANCELLED'
            """
        )

        cancelled = cursor.fetchone()["total"]

        stats = {
            "total": total,
            "registered": registered,
            "approved": approved,
            "rejected": rejected,
            "cancelled": cancelled
        }

        return render_template(
            "admin/registrations.html",
            registrations=registration_rows,
            games=games,
            total=total,
            registered=registered,
            approved=approved,
            rejected=rejected,
            cancelled=cancelled,
            stats=stats,
            search=search,
            status=status,
            selected_game_id=game_id
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Admin registrations error"
        )

        flash(
            "Unable to load registrations.",
            "error"
        )

        return redirect(
            url_for("admin.dashboard")
        )

    finally:

        close_db(
            connection,
            cursor
        )
# ============================================================
# EXPORT GAME REGISTRATIONS TO EXCEL
# ============================================================

@admin_bp.route(
    "/registrations/game/<int:game_id>/export"
)
@admin_required
def export_game_registrations(game_id):

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor()

        # ----------------------------------------------------
        # GET GAME
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT
                id,
                game_name,
                registration_mode,
                team_min_size,
                team_max_size,
                event_date,
                start_time,
                end_time,
                block,
                floor,
                room,
                prize_pool,
                status
            FROM games
            WHERE id = %s
            LIMIT 1
            """,
            (game_id,)
        )

        game = cursor.fetchone()

        if not game:
            flash(
                "Game not found.",
                "error"
            )

            return redirect(
                url_for("admin.registrations")
            )

        # ----------------------------------------------------
        # GET REGISTRATIONS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT
                r.id AS registration_id,
                r.status AS registration_status,
                r.registered_at,
                r.updated_at,

                s.student_id,
                s.name AS student_name,
                s.email,
                s.phone,
                s.branch,
                s.section,
                s.year

            FROM registrations r

            INNER JOIN students s
                ON s.id = r.student_id

            WHERE r.game_id = %s

            ORDER BY
                r.registered_at ASC,
                r.id ASC
            """,
            (game_id,)
        )

        registrations = cursor.fetchall()

        # ----------------------------------------------------
        # CREATE EXCEL WORKBOOK
        # ----------------------------------------------------

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Registrations"

        # ----------------------------------------------------
        # GAME SUMMARY
        # ----------------------------------------------------

        worksheet["A1"] = "ENGINEERS DAY 2026"

        worksheet["A2"] = "GAME"
        worksheet["B2"] = game["game_name"]

        worksheet["A3"] = "TOTAL REGISTRATIONS"
        worksheet["B3"] = len(registrations)

        worksheet["A4"] = "REGISTRATION MODE"
        worksheet["B4"] = game["registration_mode"]

        # ----------------------------------------------------
        # HEADERS
        # ----------------------------------------------------

        headers = [
            "Registration ID",
            "Registration Status",
            "Registered At",
            "Updated At",

            "Student ID",
            "Student Name",
            "Email",
            "Phone",
            "Branch",
            "Section",
            "Year",

            "Game ID",
            "Game Name",
            "Registration Mode",
            "Team Min Size",
            "Team Max Size",

            "Event Date",
            "Start Time",
            "End Time",

            "Block",
            "Floor",
            "Room",

            "Prize Pool",
            "Game Status"
        ]

        header_row = 6

        for column_number, header in enumerate(
            headers,
            start=1
        ):
            cell = worksheet.cell(
                row=header_row,
                column=column_number,
                value=header
            )

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # ----------------------------------------------------
        # DATA
        # ----------------------------------------------------

        for row_number, registration in enumerate(
            registrations,
            start=header_row + 1
        ):

            # ------------------------------------------------
            # IMPORTANT:
            # Convert timezone-aware datetime values into
            # Excel-compatible timezone-naive values.
            # ------------------------------------------------

            registered_at = registration.get(
                "registered_at"
            )

            updated_at = registration.get(
                "updated_at"
            )

            if registered_at is not None:

                try:
                    registered_at = registered_at.replace(
                        tzinfo=None
                    )
                except Exception:
                    pass

            if updated_at is not None:

                try:
                    updated_at = updated_at.replace(
                        tzinfo=None
                    )
                except Exception:
                    pass

            event_date = game.get(
                "event_date"
            )

            start_time = game.get(
                "start_time"
            )

            end_time = game.get(
                "end_time"
            )

            # ------------------------------------------------
            # DATA ROW
            # ------------------------------------------------

            values = [

                registration.get(
                    "registration_id"
                ),

                registration.get(
                    "registration_status"
                ),

                registered_at,

                updated_at,

                registration.get(
                    "student_id"
                ),

                registration.get(
                    "student_name"
                ),

                registration.get(
                    "email"
                ),

                registration.get(
                    "phone"
                ),

                registration.get(
                    "branch"
                ),

                registration.get(
                    "section"
                ),

                registration.get(
                    "year"
                ),

                game.get(
                    "id"
                ),

                game.get(
                    "game_name"
                ),

                game.get(
                    "registration_mode"
                ),

                game.get(
                    "team_min_size"
                ),

                game.get(
                    "team_max_size"
                ),

                event_date,

                start_time,

                end_time,

                game.get(
                    "block"
                ),

                game.get(
                    "floor"
                ),

                game.get(
                    "room"
                ),

                game.get(
                    "prize_pool"
                ),

                game.get(
                    "status"
                )
            ]

            for column_number, value in enumerate(
                values,
                start=1
            ):

                cell = worksheet.cell(
                    row=row_number,
                    column=column_number,
                    value=value
                )

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

            # ------------------------------------------------
            # Excel date/time formatting
            # ------------------------------------------------

            if registered_at is not None:
                worksheet.cell(
                    row=row_number,
                    column=3
                ).number_format = "yyyy-mm-dd hh:mm:ss"

            if updated_at is not None:
                worksheet.cell(
                    row=row_number,
                    column=4
                ).number_format = "yyyy-mm-dd hh:mm:ss"

            if event_date is not None:
                worksheet.cell(
                    row=row_number,
                    column=17
                ).number_format = "yyyy-mm-dd"

            if start_time is not None:
                worksheet.cell(
                    row=row_number,
                    column=18
                ).number_format = "hh:mm:ss"

            if end_time is not None:
                worksheet.cell(
                    row=row_number,
                    column=19
                ).number_format = "hh:mm:ss"

        # ----------------------------------------------------
        # COLUMN WIDTHS
        # ----------------------------------------------------

        widths = {
            "A": 18,
            "B": 22,
            "C": 24,
            "D": 24,

            "E": 18,
            "F": 26,
            "G": 32,
            "H": 18,
            "I": 18,
            "J": 12,
            "K": 10,

            "L": 12,
            "M": 30,
            "N": 20,
            "O": 16,
            "P": 16,

            "Q": 15,
            "R": 15,
            "S": 15,

            "T": 15,
            "U": 12,
            "V": 12,

            "W": 18,
            "X": 18
        }

        for column, width in widths.items():

            worksheet.column_dimensions[
                column
            ].width = width

        # ----------------------------------------------------
        # FREEZE HEADER
        # ----------------------------------------------------

        worksheet.freeze_panes = "A7"

        # ----------------------------------------------------
        # AUTO FILTER
        # ----------------------------------------------------

        if registrations:

            last_row = header_row + len(
                registrations
            )

            worksheet.auto_filter.ref = (
                f"A{header_row}:X{last_row}"
            )

        # ----------------------------------------------------
        # GAME DETAILS SHEET
        # ----------------------------------------------------

        details_sheet = workbook.create_sheet(
            "Game Details"
        )

        details = [
            ("Game ID", game.get("id")),
            ("Game Name", game.get("game_name")),
            ("Registration Mode", game.get("registration_mode")),
            ("Minimum Team Size", game.get("team_min_size")),
            ("Maximum Team Size", game.get("team_max_size")),
            ("Event Date", game.get("event_date")),
            ("Start Time", game.get("start_time")),
            ("End Time", game.get("end_time")),
            ("Block", game.get("block")),
            ("Floor", game.get("floor")),
            ("Room", game.get("room")),
            ("Prize Pool", game.get("prize_pool")),
            ("Game Status", game.get("status")),
            ("Total Registrations", len(registrations))
        ]

        for row_number, (label, value) in enumerate(
            details,
            start=1
        ):

            details_sheet.cell(
                row=row_number,
                column=1,
                value=label
            )

            details_sheet.cell(
                row=row_number,
                column=2,
                value=value
            )

            details_sheet.cell(
                row=row_number,
                column=1
            ).font = Font(
                bold=True
            )

        details_sheet.column_dimensions[
            "A"
        ].width = 28

        details_sheet.column_dimensions[
            "B"
        ].width = 40

        # ----------------------------------------------------
        # TEAM INFORMATION
        # ----------------------------------------------------

        if str(
            game.get("registration_mode") or ""
        ).upper() == "TEAM":

            team_sheet = workbook.create_sheet(
                "Team Members"
            )

            cursor.execute(
                """
                SELECT
                    t.id AS team_id,
                    t.team_name,
                    t.status AS team_status,

                    leader.student_id
                        AS leader_student_id,

                    leader.name
                        AS leader_name,

                    tm.member_role,

                    member.student_id
                        AS member_student_id,

                    member.name
                        AS member_name,

                    member.email
                        AS member_email,

                    member.phone
                        AS member_phone,

                    member.branch
                        AS member_branch,

                    member.section
                        AS member_section,

                    member.year
                        AS member_year

                FROM teams t

                INNER JOIN students leader
                    ON leader.id = t.team_leader_id

                INNER JOIN team_members tm
                    ON tm.team_id = t.id

                INNER JOIN students member
                    ON member.id = tm.student_id

                WHERE t.game_id = %s

                ORDER BY
                    t.id ASC,
                    tm.id ASC
                """,
                (game_id,)
            )

            team_members = cursor.fetchall()

            team_headers = [
                "Team ID",
                "Team Name",
                "Team Status",
                "Leader Student ID",
                "Leader Name",
                "Member Role",
                "Member Student ID",
                "Member Name",
                "Member Email",
                "Member Phone",
                "Branch",
                "Section",
                "Year"
            ]

            for column_number, header in enumerate(
                team_headers,
                start=1
            ):

                cell = team_sheet.cell(
                    row=1,
                    column=column_number,
                    value=header
                )

                cell.font = Font(
                    bold=True
                )

            for row_number, member in enumerate(
                team_members,
                start=2
            ):

                team_values = [

                    member.get(
                        "team_id"
                    ),

                    member.get(
                        "team_name"
                    ),

                    member.get(
                        "team_status"
                    ),

                    member.get(
                        "leader_student_id"
                    ),

                    member.get(
                        "leader_name"
                    ),

                    member.get(
                        "member_role"
                    ),

                    member.get(
                        "member_student_id"
                    ),

                    member.get(
                        "member_name"
                    ),

                    member.get(
                        "member_email"
                    ),

                    member.get(
                        "member_phone"
                    ),

                    member.get(
                        "member_branch"
                    ),

                    member.get(
                        "member_section"
                    ),

                    member.get(
                        "member_year"
                    )
                ]

                for column_number, value in enumerate(
                    team_values,
                    start=1
                ):

                    team_sheet.cell(
                        row=row_number,
                        column=column_number,
                        value=value
                    )

            team_widths = {
                "A": 12,
                "B": 24,
                "C": 18,
                "D": 20,
                "E": 25,
                "F": 16,
                "G": 20,
                "H": 25,
                "I": 32,
                "J": 18,
                "K": 18,
                "L": 12,
                "M": 10
            }

            for column, width in team_widths.items():

                team_sheet.column_dimensions[
                    column
                ].width = width

            team_sheet.freeze_panes = "A2"

        # ----------------------------------------------------
        # SAVE WORKBOOK TO MEMORY
        # ----------------------------------------------------

        output = io.BytesIO()

        workbook.save(output)

        output.seek(0)

        # ----------------------------------------------------
        # SAFE FILE NAME
        # ----------------------------------------------------

        safe_game_name = re.sub(
            r"[^A-Za-z0-9_-]+",
            "_",
            str(
                game.get("game_name") or
                "Game"
            )
        ).strip("_")

        if not safe_game_name:
            safe_game_name = "Game"

        filename = (
            f"{safe_game_name}_Registrations.xlsx"
        )

        # ----------------------------------------------------
        # SEND FILE
        # ----------------------------------------------------

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    except Exception:

        if connection:
            try:
                connection.rollback()
            except Exception:
                pass

        current_app.logger.exception(
            "Game registration Excel export error"
        )

        flash(
            "Unable to export registrations.",
            "admin_error"
        )

        return redirect(
            url_for("admin.registrations")
        )

    finally:

        if cursor:
            try:
                cursor.close()
            except Exception:
                pass

        if connection:
            try:
                connection.close()
            except Exception:
                pass
# ============================================================
# REGISTRATION DETAILS
# ============================================================

@admin_bp.route("/registrations/<int:registration_id>")
@admin_required
def registration_details(registration_id):

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor()

        # ----------------------------------------------------
        # 1. REGISTRATION + STUDENT + GAME
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                r.id,
                r.student_id,
                r.game_id,
                r.status,
                r.registered_at,
                r.updated_at,

                -- Student
                s.name AS name,
                s.email AS email,
                s.phone AS phone,
                s.student_id AS student_roll,
                s.branch,
                s.section,
                s.year,

                -- Game
                g.game_name,
                g.description,
                g.registration_mode,
                g.team_min_size,
                g.team_max_size,
                g.event_date AS schedule_date,
                g.start_time AS schedule_time,
                CONCAT_WS(', ', g.block, g.floor, g.room) AS venue,
                g.prize_pool,
                g.rules,
                g.requirements,
                g.registration_open,
                g.status AS game_status

            FROM registrations r

            INNER JOIN students s
                ON s.id = r.student_id

            INNER JOIN games g
                ON g.id = r.game_id

            WHERE r.id = %s

            LIMIT 1
        """, (registration_id,))

        registration = cursor.fetchone()

        if not registration:

            flash(
                "Registration not found.",
                "error"
            )

            return redirect(
                url_for("admin.registrations")
            )

        # ----------------------------------------------------
        # 2. DEFAULT TEAM VALUES
        # ----------------------------------------------------

        team_members = []

        registration["team_id"] = None
        registration["team_name"] = None
        registration["team_status"] = None
        registration["team_leader_name"] = None

        # ----------------------------------------------------
        # 3. FIND TEAM FOR THIS REGISTRATION
        #
        # A student can be connected to a team through
        # team_members.
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                t.id AS team_id,
                t.team_name,
                t.status AS team_status,
                t.team_leader_id,

                leader.name AS team_leader_name

            FROM team_members tm

            INNER JOIN teams t
                ON t.id = tm.team_id

            INNER JOIN students leader
                ON leader.id = t.team_leader_id

            WHERE
                tm.student_id = %s
                AND t.game_id = %s

            ORDER BY t.id DESC

            LIMIT 1
        """, (
            registration["student_id"],
            registration["game_id"]
        ))

        team = cursor.fetchone()

        # ----------------------------------------------------
        # 4. TEAM DATA
        # ----------------------------------------------------

        if team:

            registration["team_id"] = team["team_id"]
            registration["team_name"] = team["team_name"]
            registration["team_status"] = team["team_status"]
            registration["team_leader_name"] = team["team_leader_name"]

            # ------------------------------------------------
            # 5. LOAD ALL TEAM MEMBERS
            # ------------------------------------------------

            cursor.execute("""
                SELECT

                    tm.id AS team_member_id,

                    tm.member_role,
                    tm.joined_at,

                    s.id AS student_id,
                    s.student_id AS student_roll,
                    s.name,
                    s.email,
                    s.phone,
                    s.branch,
                    s.section,
                    s.year

                FROM team_members tm

                INNER JOIN students s
                    ON s.id = tm.student_id

                WHERE tm.team_id = %s

                ORDER BY
                    CASE
                        WHEN tm.member_role = 'LEADER'
                        THEN 0
                        ELSE 1
                    END,
                    tm.id ASC
            """, (team["team_id"],))

            team_members = cursor.fetchall()

        # ----------------------------------------------------
        # 6. RENDER DETAILS PAGE
        # ----------------------------------------------------

        return render_template(
            "admin/registration_details.html",
            registration=registration,
            team_members=team_members
        )

    except Exception:

        if connection:
            connection.rollback()

        current_app.logger.exception(
            "Registration details error"
        )

        flash(
            "Unable to load registration details.",
            "error"
        )

        return redirect(
            url_for("admin.registrations")
        )

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ============================================================
# REGISTRATION STATUS UPDATE
# ============================================================

@admin_bp.route(
    "/registrations/<int:registration_id>/status",
    methods=["POST"]
)
@admin_required
def update_registration_status(
    registration_id
):

    new_status = request.form.get(
        "status",
        ""
    ).strip().upper()

    if new_status not in ALLOWED_REGISTRATION_STATUSES:

        flash(
            "Invalid registration status.",
            "error"
        )

        return redirect(
            url_for(
                "admin.registration_details",
                registration_id=registration_id
            )
        )

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT id
            FROM registrations
            WHERE id = %s
            LIMIT 1
            """,
            (registration_id,)
        )

        if not cursor.fetchone():

            flash(
                "Registration not found.",
                "error"
            )

            return redirect(
                url_for("admin.registrations")
            )

        cursor.execute(
            """
            UPDATE registrations
            SET
                status = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
            """,
            (
                new_status,
                registration_id
            )
        )

        connection.commit()

        flash(
            f"Registration status changed to {new_status}.",
            "success"
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Registration status update error"
        )

        flash(
            "Unable to update registration status.",
            "error"
        )

    finally:

        close_db(
            connection,
            cursor
        )

    return redirect(
        url_for(
            "admin.registration_details",
            registration_id=registration_id
        )
    )


# ============================================================
# RESULTS MANAGEMENT
# ============================================================

@admin_bp.route("/results")
@admin_required
def results():

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        selected_game_id = request.args.get(
            "game_id",
            ""
        ).strip()

        selected_status = request.args.get(
            "status",
            ""
        ).strip().lower()

        query = """
            SELECT

                g.id AS game_id,
                g.game_name,
                g.registration_mode,
                g.event_date,
                g.start_time,
                g.end_time,
                g.block,
                g.floor,
                g.room,
                g.status AS game_status,

                r.id AS result_id,
                r.winner,
                r.runner_up,
                r.winner_prize,
                r.runner_up_prize,
                r.winner_certificate,
                r.runner_up_certificate,
                r.result_details,
                r.is_published,
                r.published_at,
                r.created_at

            FROM games g

            LEFT JOIN results r
                ON r.game_id = g.id
        """

        conditions = []
        params = []

        # ----------------------------------------------------
        # EVENT FILTER
        # ----------------------------------------------------

        if selected_game_id:

            game_id_value = clean_int(
                selected_game_id
            )

            if game_id_value is not None:

                conditions.append(
                    "g.id = %s"
                )

                params.append(
                    game_id_value
                )

            else:

                selected_game_id = ""

        # ----------------------------------------------------
        # RESULT STATUS FILTER
        # ----------------------------------------------------

        if selected_status == "published":

            conditions.append(
                "r.id IS NOT NULL AND r.is_published = TRUE"
            )

        elif selected_status == "pending":

            conditions.append(
                "r.id IS NULL"
            )

        elif selected_status == "draft":

            conditions.append(
                "r.id IS NOT NULL AND r.is_published = FALSE"
            )

        # ----------------------------------------------------
        # WHERE
        # ----------------------------------------------------

        if conditions:

            query += (
                " WHERE "
                + " AND ".join(conditions)
            )

        # ----------------------------------------------------
        # ORDER
        # ----------------------------------------------------

        query += """
            ORDER BY

                CASE
                    WHEN g.event_date IS NULL THEN 1
                    ELSE 0
                END ASC,

                g.event_date ASC,

                CASE
                    WHEN g.start_time IS NULL THEN 1
                    ELSE 0
                END ASC,

                g.start_time ASC,

                g.id ASC
        """

        cursor.execute(
            query,
            tuple(params)
        )

        result_rows = cursor.fetchall()

        # ----------------------------------------------------
        # EVENTS FOR FILTER
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT
                id,
                game_name
            FROM games
            ORDER BY game_name ASC
            """
        )

        games = cursor.fetchall()

        # ----------------------------------------------------
        # TOTAL RESULTS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM results
            """
        )

        total_results = cursor.fetchone()["total"]

        # ----------------------------------------------------
        # EVENTS WITH RESULTS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(DISTINCT game_id) AS total
            FROM results
            """
        )

        completed_events = cursor.fetchone()["total"]

        # ----------------------------------------------------
        # EVENTS WITHOUT RESULTS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM games g
            LEFT JOIN results r
                ON r.game_id = g.id
            WHERE r.id IS NULL
            """
        )

        pending_events = cursor.fetchone()["total"]

        # ----------------------------------------------------
        # TOTAL EVENTS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM games
            """
        )

        total_events = cursor.fetchone()["total"]

        return render_template(
            "admin/results.html",

            results=result_rows,
            games=games,

            total_results=total_results,
            completed_events=completed_events,
            pending_events=pending_events,
            total_events=total_events,

            selected_game_id=selected_game_id,
            selected_status=selected_status
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Admin results error"
        )

        flash(
            "Unable to load results management.",
            "error"
        )

        return redirect(
            url_for("admin.dashboard")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# CREATE RESULT
# ============================================================

@admin_bp.route(
    "/results/create",
    methods=["GET", "POST"]
)
@admin_required
def result_create():

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        # ----------------------------------------------------
        # POST
        # ----------------------------------------------------

        if request.method == "POST":

            raw_game_id = request.form.get(
                "game_id",
                ""
            ).strip()

            game_id = clean_int(
                raw_game_id
            )

            if game_id is None:

                flash(
                    "Please select a valid event.",
                    "error"
                )

                return redirect(
                    url_for("admin.result_create")
                )

            # ------------------------------------------------
            # GET GAME
            # ------------------------------------------------

            cursor.execute(
                """
                SELECT
                    id,
                    game_name,
                    registration_mode
                FROM games
                WHERE id = %s
                LIMIT 1
                """,
                (game_id,)
            )

            game = cursor.fetchone()

            if not game:

                flash(
                    "Event not found.",
                    "error"
                )

                return redirect(
                    url_for("admin.result_create")
                )

            # ------------------------------------------------
            # FORM DATA
            # ------------------------------------------------

            winner = (
                request.form.get(
                    "winner",
                    ""
                ).strip()
                or None
            )

            runner_up = (
                request.form.get(
                    "runner_up",
                    ""
                ).strip()
                or None
            )

            winner_student_id = clean_int(
                request.form.get(
                    "winner_student_id"
                )
            )

            winner_team_id = clean_int(
                request.form.get(
                    "winner_team_id"
                )
            )

            runner_up_student_id = clean_int(
                request.form.get(
                    "runner_up_student_id"
                )
            )

            runner_up_team_id = clean_int(
                request.form.get(
                    "runner_up_team_id"
                )
            )

            winner_prize = (
                request.form.get(
                    "winner_prize",
                    ""
                ).strip()
                or None
            )

            runner_up_prize = (
                request.form.get(
                    "runner_up_prize",
                    ""
                ).strip()
                or None
            )

            winner_certificate = form_checkbox(
                "winner_certificate"
            )

            runner_up_certificate = form_checkbox(
                "runner_up_certificate"
            )

            result_details = (
                request.form.get(
                    "result_details",
                    ""
                ).strip()
                or None
            )

            is_published = form_checkbox(
                "is_published"
            )

            # ------------------------------------------------
            # CHECK EXISTING RESULT
            # ------------------------------------------------

            cursor.execute(
                """
                SELECT id
                FROM results
                WHERE game_id = %s
                LIMIT 1
                """,
                (game_id,)
            )

            existing = cursor.fetchone()

            if existing:

                flash(
                    "A result already exists for this event. "
                    "Use Edit instead.",
                    "error"
                )

                return redirect(
                    url_for(
                        "admin.result_edit",
                        game_id=game_id
                    )
                )

            # ------------------------------------------------
            # INSERT RESULT
            #
            # PostgreSQL:
            # TRUE/FALSE
            # CURRENT_TIMESTAMP / NOW()
            # ------------------------------------------------

            if is_published:

                cursor.execute(
                    """
                    INSERT INTO results
                    (
                        game_id,
                        winner,
                        winner_student_id,
                        winner_team_id,
                        runner_up,
                        runner_up_student_id,
                        runner_up_team_id,
                        winner_prize,
                        runner_up_prize,
                        winner_certificate,
                        runner_up_certificate,
                        result_details,
                        is_published,
                        published_at,
                        published_by
                    )
                    VALUES
                    (
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        TRUE,
                        CURRENT_TIMESTAMP,
                        %s
                    )
                    """,
                    (
                        game_id,
                        winner,
                        winner_student_id,
                        winner_team_id,
                        runner_up,
                        runner_up_student_id,
                        runner_up_team_id,
                        winner_prize,
                        runner_up_prize,
                        winner_certificate,
                        runner_up_certificate,
                        result_details,
                        session.get("admin_id")
                    )
                )

            else:

                cursor.execute(
                    """
                    INSERT INTO results
                    (
                        game_id,
                        winner,
                        winner_student_id,
                        winner_team_id,
                        runner_up,
                        runner_up_student_id,
                        runner_up_team_id,
                        winner_prize,
                        runner_up_prize,
                        winner_certificate,
                        runner_up_certificate,
                        result_details,
                        is_published,
                        published_at,
                        published_by
                    )
                    VALUES
                    (
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        FALSE,
                        NULL,
                        NULL
                    )
                    """,
                    (
                        game_id,
                        winner,
                        winner_student_id,
                        winner_team_id,
                        runner_up,
                        runner_up_student_id,
                        runner_up_team_id,
                        winner_prize,
                        runner_up_prize,
                        winner_certificate,
                        runner_up_certificate,
                        result_details
                    )
                )

            connection.commit()

            if is_published:

                flash(
                    "Result published successfully.",
                    "success"
                )

            else:

                flash(
                    "Result saved as draft.",
                    "success"
                )

            return redirect(
                url_for("admin.results")
            )

        # ----------------------------------------------------
        # GET
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT
                g.id,
                g.game_name,
                g.registration_mode,
                g.event_date
            FROM games g

            LEFT JOIN results r
                ON r.game_id = g.id

            WHERE r.id IS NULL

            ORDER BY
                CASE
                    WHEN g.event_date IS NULL THEN 1
                    ELSE 0
                END,
                g.event_date ASC,
                g.game_name ASC
            """
        )

        games = cursor.fetchall()

        return render_template(
            "admin/result_edit.html",
            game=None,
            result=None,
            games=games,
            mode="create"
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Result creation error"
        )

        flash(
            "Unable to create result.",
            "error"
        )

        return redirect(
            url_for("admin.results")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# EDIT RESULT
# ============================================================

@admin_bp.route(
    "/results/<int:game_id>/edit",
    methods=["GET", "POST"]
)
@admin_required
def result_edit(game_id):

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        # ----------------------------------------------------
        # LOAD EVENT
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT
                id,
                game_name,
                registration_mode,
                event_date
            FROM games
            WHERE id = %s
            LIMIT 1
            """,
            (game_id,)
        )

        game = cursor.fetchone()

        if not game:

            flash(
                "Event not found.",
                "error"
            )

            return redirect(
                url_for("admin.results")
            )

        # ----------------------------------------------------
        # POST
        # ----------------------------------------------------

        if request.method == "POST":

            winner = (
                request.form.get(
                    "winner",
                    ""
                ).strip()
                or None
            )

            runner_up = (
                request.form.get(
                    "runner_up",
                    ""
                ).strip()
                or None
            )

            winner_student_id = clean_int(
                request.form.get(
                    "winner_student_id"
                )
            )

            winner_team_id = clean_int(
                request.form.get(
                    "winner_team_id"
                )
            )

            runner_up_student_id = clean_int(
                request.form.get(
                    "runner_up_student_id"
                )
            )

            runner_up_team_id = clean_int(
                request.form.get(
                    "runner_up_team_id"
                )
            )

            winner_prize = (
                request.form.get(
                    "winner_prize",
                    ""
                ).strip()
                or None
            )

            runner_up_prize = (
                request.form.get(
                    "runner_up_prize",
                    ""
                ).strip()
                or None
            )

            winner_certificate = form_checkbox(
                "winner_certificate"
            )

            runner_up_certificate = form_checkbox(
                "runner_up_certificate"
            )

            result_details = (
                request.form.get(
                    "result_details",
                    ""
                ).strip()
                or None
            )

            is_published = form_checkbox(
                "is_published"
            )

            # ------------------------------------------------
            # CHECK RESULT
            # ------------------------------------------------

            cursor.execute(
                """
                SELECT id
                FROM results
                WHERE game_id = %s
                LIMIT 1
                """,
                (game_id,)
            )

            existing_result = cursor.fetchone()

            if not existing_result:

                flash(
                    "No result exists for this event.",
                    "error"
                )

                return redirect(
                    url_for("admin.results")
                )

            # ------------------------------------------------
            # PUBLISHED
            # ------------------------------------------------

            if is_published:

                cursor.execute(
                    """
                    UPDATE results
                    SET
                        winner = %s,
                        winner_student_id = %s,
                        winner_team_id = %s,

                        runner_up = %s,
                        runner_up_student_id = %s,
                        runner_up_team_id = %s,

                        winner_prize = %s,
                        runner_up_prize = %s,

                        winner_certificate = %s,
                        runner_up_certificate = %s,

                        result_details = %s,

                        is_published = TRUE,

                        published_at = COALESCE(
                            published_at,
                            CURRENT_TIMESTAMP
                        ),

                        published_by = %s

                    WHERE game_id = %s
                    """,
                    (
                        winner,
                        winner_student_id,
                        winner_team_id,

                        runner_up,
                        runner_up_student_id,
                        runner_up_team_id,

                        winner_prize,
                        runner_up_prize,

                        winner_certificate,
                        runner_up_certificate,

                        result_details,

                        session.get(
                            "admin_id"
                        ),

                        game_id
                    )
                )

            # ------------------------------------------------
            # DRAFT
            # ------------------------------------------------

            else:

                cursor.execute(
                    """
                    UPDATE results
                    SET
                        winner = %s,
                        winner_student_id = %s,
                        winner_team_id = %s,

                        runner_up = %s,
                        runner_up_student_id = %s,
                        runner_up_team_id = %s,

                        winner_prize = %s,
                        runner_up_prize = %s,

                        winner_certificate = %s,
                        runner_up_certificate = %s,

                        result_details = %s,

                        is_published = FALSE,
                        published_at = NULL,
                        published_by = NULL

                    WHERE game_id = %s
                    """,
                    (
                        winner,
                        winner_student_id,
                        winner_team_id,

                        runner_up,
                        runner_up_student_id,
                        runner_up_team_id,

                        winner_prize,
                        runner_up_prize,

                        winner_certificate,
                        runner_up_certificate,

                        result_details,

                        game_id
                    )
                )

            connection.commit()

            if is_published:

                flash(
                    "Result published successfully.",
                    "success"
                )

            else:

                flash(
                    "Result saved as draft.",
                    "success"
                )

            return redirect(
                url_for("admin.results")
            )

        # ----------------------------------------------------
        # GET RESULT
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT *
            FROM results
            WHERE game_id = %s
            LIMIT 1
            """,
            (game_id,)
        )

        result = cursor.fetchone()

        if not result:

            flash(
                "No result exists for this event yet.",
                "error"
            )

            return redirect(
                url_for(
                    "admin.result_create",
                    game_id=game_id
                )
            )

        return render_template(
            "admin/result_edit.html",
            game=game,
            result=result,
            games=None,
            mode="edit"
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Result edit error"
        )

        flash(
            "Unable to update result.",
            "error"
        )

        return redirect(
            url_for("admin.results")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# NOTIFICATION MANAGEMENT
# ============================================================

@admin_bp.route("/notifications")
@admin_required
def notifications():

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        search = request.args.get(
            "q",
            ""
        ).strip()

        notification_type = request.args.get(
            "type",
            "all"
        ).strip().upper()

        query = """
            SELECT
                n.id,
                n.game_id,
                n.title,
                n.message,
                n.notification_type,
                n.is_published,
                n.published_at,
                n.expires_at,
                n.created_by,
                n.created_at,
                n.updated_at,

                g.game_name

            FROM notifications n

            LEFT JOIN games g
                ON g.id = n.game_id
        """

        conditions = []
        params = []

        # ----------------------------------------------------
        # SEARCH
        # ----------------------------------------------------

        if search:

            value = f"%{search}%"

            conditions.append(
                """
                (
                    n.title ILIKE %s
                    OR n.message ILIKE %s
                    OR g.game_name ILIKE %s
                )
                """
            )

            params.extend([
                value,
                value,
                value
            ])

        # ----------------------------------------------------
        # TYPE
        # ----------------------------------------------------

        if notification_type in ALLOWED_NOTIFICATION_TYPES:

            conditions.append(
                "n.notification_type = %s"
            )

            params.append(
                notification_type
            )

        if conditions:

            query += (
                " WHERE "
                + " AND ".join(conditions)
            )

        query += """
            ORDER BY
                n.created_at DESC,
                n.id DESC
            LIMIT 1000
        """

        cursor.execute(
            query,
            tuple(params)
        )

        notification_rows = cursor.fetchall()

        # ----------------------------------------------------
        # EVENTS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT
                id,
                game_name
            FROM games
            ORDER BY game_name ASC
            """
        )

        games = cursor.fetchall()

        return render_template(
            "admin/notifications.html",
            notifications=notification_rows,
            games=games,
            search=search,
            notification_type=notification_type
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Admin notifications error"
        )

        flash(
            "Unable to load notifications.",
            "error"
        )

        return redirect(
            url_for("admin.dashboard")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# CREATE NOTIFICATION
# ============================================================

@admin_bp.route(
    "/notifications/create",
    methods=["GET", "POST"]
)
@admin_required
def notification_create():

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        if request.method == "POST":

            title = request.form.get(
                "title",
                ""
            ).strip()

            message = request.form.get(
                "message",
                ""
            ).strip()

            notification_type = request.form.get(
                "notification_type",
                "GENERAL"
            ).strip().upper()

            raw_game_id = request.form.get(
                "game_id",
                ""
            ).strip()

            game_id = (
                int(raw_game_id)
                if raw_game_id.isdigit()
                else None
            )

            expires_at = request.form.get(
                "expires_at",
                ""
            ).strip() or None

            is_published = form_checkbox(
                "is_published"
            )

            # ------------------------------------------------
            # VALIDATION
            # ------------------------------------------------

            if not title:

                flash(
                    "Notification title is required.",
                    "error"
                )

                return redirect(
                    url_for(
                        "admin.notification_create"
                    )
                )

            if not message:

                flash(
                    "Notification message is required.",
                    "error"
                )

                return redirect(
                    url_for(
                        "admin.notification_create"
                    )
                )

            if notification_type not in ALLOWED_NOTIFICATION_TYPES:

                flash(
                    "Invalid notification type.",
                    "error"
                )

                return redirect(
                    url_for(
                        "admin.notification_create"
                    )
                )

            # ------------------------------------------------
            # OPTIONAL GAME VALIDATION
            # ------------------------------------------------

            if game_id is not None:

                cursor.execute(
                    """
                    SELECT id
                    FROM games
                    WHERE id = %s
                    LIMIT 1
                    """,
                    (game_id,)
                )

                if not cursor.fetchone():

                    flash(
                        "Selected event was not found.",
                        "error"
                    )

                    return redirect(
                        url_for(
                            "admin.notification_create"
                        )
                    )

            # ------------------------------------------------
            # INSERT
            # ------------------------------------------------

            if is_published:

                cursor.execute(
                    """
                    INSERT INTO notifications
                    (
                        game_id,
                        title,
                        message,
                        notification_type,
                        is_published,
                        published_at,
                        expires_at,
                        created_by
                    )
                    VALUES
                    (
                        %s,
                        %s,
                        %s,
                        %s,
                        TRUE,
                        CURRENT_TIMESTAMP,
                        %s,
                        %s
                    )
                    """,
                    (
                        game_id,
                        title,
                        message,
                        notification_type,
                        expires_at,
                        session.get("admin_id")
                    )
                )

            else:

                cursor.execute(
                    """
                    INSERT INTO notifications
                    (
                        game_id,
                        title,
                        message,
                        notification_type,
                        is_published,
                        published_at,
                        expires_at,
                        created_by
                    )
                    VALUES
                    (
                        %s,
                        %s,
                        %s,
                        %s,
                        FALSE,
                        NULL,
                        %s,
                        %s
                    )
                    """,
                    (
                        game_id,
                        title,
                        message,
                        notification_type,
                        expires_at,
                        session.get("admin_id")
                    )
                )

            connection.commit()

            flash(
                "Notification created successfully.",
                "success"
            )

            return redirect(
                url_for("admin.notifications")
            )

        # ----------------------------------------------------
        # GET
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT
                id,
                game_name
            FROM games
            ORDER BY game_name ASC
            """
        )

        games = cursor.fetchall()

        return render_template(
            "admin/notifications.html",
            notifications=[],
            games=games,
            search="",
            notification_type="all",
            create_mode=True
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Notification creation error"
        )

        flash(
            "Unable to create notification.",
            "error"
        )

        return redirect(
            url_for("admin.notifications")
        )

    finally:

        close_db(
            connection,
            cursor
        )


# ============================================================
# TOGGLE NOTIFICATION
# ============================================================

@admin_bp.route(
    "/notifications/<int:notification_id>/toggle",
    methods=["POST"]
)
@admin_required
def notification_toggle(
    notification_id
):

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT
                id,
                title,
                is_published
            FROM notifications
            WHERE id = %s
            LIMIT 1
            """,
            (notification_id,)
        )

        notification = cursor.fetchone()

        if not notification:

            flash(
                "Notification not found.",
                "error"
            )

            return redirect(
                url_for("admin.notifications")
            )

        # ----------------------------------------------------
        # UNPUBLISH
        # ----------------------------------------------------

        if bool(notification["is_published"]):

            cursor.execute(
                """
                UPDATE notifications
                SET
                    is_published = FALSE,
                    published_at = NULL,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = %s
                """,
                (notification_id,)
            )

            message = (
                "Notification unpublished."
            )

        # ----------------------------------------------------
        # PUBLISH
        # ----------------------------------------------------

        else:

            cursor.execute(
                """
                UPDATE notifications
                SET
                    is_published = TRUE,
                    published_at = CURRENT_TIMESTAMP,
                    created_by = COALESCE(
                        created_by,
                        %s
                    ),
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = %s
                """,
                (
                    session.get("admin_id"),
                    notification_id
                )
            )

            message = (
                "Notification published."
            )

        connection.commit()

        flash(
            message,
            "success"
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Notification toggle error"
        )

        flash(
            "Unable to change notification status.",
            "error"
        )

    finally:

        close_db(
            connection,
            cursor
        )

    return redirect(
        url_for("admin.notifications")
    )


# ============================================================
# DELETE NOTIFICATION
# ============================================================

@admin_bp.route(
    "/notifications/<int:notification_id>/delete",
    methods=["POST"]
)
@admin_required
def notification_delete(
    notification_id
):

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT id
            FROM notifications
            WHERE id = %s
            LIMIT 1
            """,
            (notification_id,)
        )

        if not cursor.fetchone():

            flash(
                "Notification not found.",
                "error"
            )

            return redirect(
                url_for("admin.notifications")
            )

        cursor.execute(
            """
            DELETE FROM notifications
            WHERE id = %s
            """,
            (notification_id,)
        )

        connection.commit()

        flash(
            "Notification deleted successfully.",
            "success"
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Notification deletion error"
        )

        flash(
            "Unable to delete notification.",
            "error"
        )

    finally:

        close_db(
            connection,
            cursor
        )

    return redirect(
        url_for("admin.notifications")
    )


# ============================================================
# ADMIN SETTINGS
#
# IMPORTANT:
# The portal_settings table is already created by the
# Supabase PostgreSQL schema.
#
# We DO NOT create the table inside a web request.
# ============================================================

@admin_bp.route(
    "/settings",
    methods=["GET", "POST"]
)
@admin_required
def settings():

    connection = None
    cursor = None

    try:

        connection = get_db_connection()
        cursor = connection.cursor()

        # ----------------------------------------------------
        # ENSURE DEFAULT SETTINGS ROW EXISTS
        #
        # PostgreSQL:
        # ON CONFLICT DO NOTHING
        # ----------------------------------------------------

        cursor.execute(
            """
            INSERT INTO portal_settings
            (
                id,
                portal_name,
                event_date,
                portal_status,
                registration_enabled,
                student_login_enabled
            )
            VALUES
            (
                1,
                'Engineers Day 2026',
                '2026-09-15',
                'ACTIVE',
                TRUE,
                TRUE
            )
            ON CONFLICT (id)
            DO NOTHING
            """
        )

        connection.commit()

        # ----------------------------------------------------
        # SAVE SETTINGS
        # ----------------------------------------------------

        if request.method == "POST":

            portal_name = request.form.get(
                "portal_name",
                ""
            ).strip()

            event_date = request.form.get(
                "event_date",
                ""
            ).strip() or None

            portal_status = request.form.get(
                "portal_status",
                "ACTIVE"
            ).strip().upper()

            registration_enabled = form_checkbox(
                "registration_enabled"
            )

            student_login_enabled = form_checkbox(
                "student_login_enabled"
            )

            # ------------------------------------------------
            # VALIDATION
            # ------------------------------------------------

            if not portal_name:

                flash(
                    "Portal name is required.",
                    "error"
                )

                return redirect(
                    url_for("admin.settings")
                )

            if portal_status not in ALLOWED_PORTAL_STATUSES:

                flash(
                    "Invalid portal status.",
                    "error"
                )

                return redirect(
                    url_for("admin.settings")
                )

            # ------------------------------------------------
            # UPDATE
            #
            # PostgreSQL boolean values are Python True/False.
            # updated_at is explicitly maintained here.
            # ------------------------------------------------

            cursor.execute(
                """
                UPDATE portal_settings
                SET
                    portal_name = %s,
                    event_date = %s,
                    portal_status = %s,
                    registration_enabled = %s,
                    student_login_enabled = %s,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = 1
                """,
                (
                    portal_name,
                    event_date,
                    portal_status,
                    registration_enabled,
                    student_login_enabled
                )
            )

            connection.commit()

            flash(
                "Portal settings updated successfully.",
                "success"
            )

            return redirect(
                url_for("admin.settings")
            )

        # ----------------------------------------------------
        # LOAD SETTINGS
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT
                id,
                portal_name,
                event_date,
                portal_status,
                registration_enabled,
                student_login_enabled,
                updated_at
            FROM portal_settings
            WHERE id = 1
            LIMIT 1
            """
        )

        settings_data = cursor.fetchone()

        # ----------------------------------------------------
        # EVENT COUNT
        # ----------------------------------------------------

        cursor.execute(
            """
            SELECT COUNT(*) AS total
            FROM games
            """
        )

        total_events = cursor.fetchone()["total"]

        return render_template(
            "admin/settings.html",
            settings=settings_data,
            total_events=total_events
        )

    except Exception:

        rollback_db(connection)

        current_app.logger.exception(
            "Admin settings error"
        )

        flash(
            "Unable to load settings.",
            "error"
        )

        return redirect(
            url_for("admin.dashboard")
        )

    finally:

        close_db(
            connection,
            cursor
        )